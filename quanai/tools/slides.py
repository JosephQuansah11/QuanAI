"""PowerPoint generation - real .pptx files, 16:9, themed and readable."""

from __future__ import annotations

from pathlib import Path

from ..errors import MissingDependency, ToolError
from .registry import ToolContext, ToolResult, artifact

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Light is the default. Dark palettes stay available for people who ask for them.
THEMES: dict[str, dict[str, str]] = {
    "paper": {
        "bg": "FAFAF7", "fg": "1F2937", "accent": "B45309",
        "muted": "6B7280", "panel": "F1EFE8",
    },
    "sky": {
        "bg": "F5F9FF", "fg": "0F172A", "accent": "2563EB",
        "muted": "64748B", "panel": "E7F0FB",
    },
    "mint": {
        "bg": "F4FAF7", "fg": "10352C", "accent": "047857",
        "muted": "4F6F66", "panel": "E4F3EC",
    },
    "sand": {
        "bg": "FDF8F1", "fg": "3B2F2A", "accent": "C2410C",
        "muted": "7A6A61", "panel": "F6EADC",
    },
    # opt-in dark palettes
    "midnight": {
        "bg": "0F172A", "fg": "F8FAFC", "accent": "38BDF8",
        "muted": "94A3B8", "panel": "1E293B",
    },
    "forest": {
        "bg": "0B2B26", "fg": "ECFDF5", "accent": "34D399",
        "muted": "9CBFB4", "panel": "123B33",
    },
    "plum": {
        "bg": "2A0E2E", "fg": "FDF4FF", "accent": "F472B6",
        "muted": "C4A5C9", "panel": "3B1540",
    },
}

LIGHT_THEMES = ("paper", "sky", "mint", "sand")
DEFAULT_THEME = "paper"


WIDTH = 13.333
HEIGHT = 7.5


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor

    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _rect(slide, left, top, width, height, color: str):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _textbox(
    slide, left, top, width, height, text, size, color,
    bold=False, align="left", spacing=1.0, font="Calibri", anchor="top",
):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP if anchor == "top" else MSO_ANCHOR.MIDDLE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = alignment
        paragraph.line_spacing = spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = font
            run.font.color.rgb = _rgb(color)
    return box


def _bullets(slide, items, theme, top=1.75, size=20):
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.6), Inches(HEIGHT - top - 0.9))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_top = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.line_spacing = 1.15
        paragraph.space_after = Pt(12)
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.name = "Calibri"
            run.font.color.rgb = _rgb(theme["fg"])
    return box


def _accent_rule(slide, theme, left=0.9, top=1.45, width=1.6):
    _rect(slide, left, top, width, 0.055, theme["accent"])


def _blank(prs, theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = _rgb(theme["bg"])
    return slide


def _fit(text, limit: int) -> str:
    flat = " ".join(str(text).split())
    return flat if len(flat) <= limit else flat[: limit - 1].rstrip() + "…"



def build_deck(
    path: Path,
    title: str,
    slides: list[dict],
    subtitle: str = "",
    theme_name: str = DEFAULT_THEME,
    author: str = "QuanAI",
) -> dict:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.util import Inches
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise MissingDependency("slides.create", "python-pptx", "pip install python-pptx") from exc

    if not slides:
        raise ToolError("a deck needs at least one slide")

    theme = THEMES.get(theme_name) or THEMES[DEFAULT_THEME]
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)

    for index, raw in enumerate(slides):
        kind = str(raw.get("kind") or "content").lower()
        slide_title = _fit(raw.get("title", ""), 80)
        bullets = [str(b) for b in (raw.get("bullets") or []) if str(b).strip()]
        notes = str(raw.get("notes") or "")

        if kind == "title" or index == 0:
            slide = _blank(prs, theme)
            _rect(slide, 0, 0, 0.22, HEIGHT, theme["accent"])
            _textbox(slide, 1.1, 2.35, 10.9, 1.9, _fit(slide_title, 70), 46, theme["fg"], bold=True)
            _accent_rule(slide, theme, left=1.1, top=4.35, width=2.2)
            sub = subtitle or (bullets[0] if bullets else "")
            if sub:
                _textbox(slide, 1.1, 4.65, 10.6, 1.2, _fit(sub, 190), 20, theme["muted"])
            if author:
                _textbox(slide, 1.1, 6.55, 6.0, 0.5, author, 12, theme["muted"])
        elif kind == "image" and raw.get("image"):
            slide = _blank(prs, theme)
            image_path = Path(str(raw["image"]))
            if image_path.exists():
                try:
                    slide.shapes.add_picture(
                        str(image_path), Inches(0), Inches(0),
                        width=Inches(WIDTH), height=Inches(HEIGHT),
                    )
                except Exception:
                    pass
            _rect(slide, 0, HEIGHT - 2.1, WIDTH, 2.1, theme["bg"])
            _textbox(slide, 0.9, HEIGHT - 1.75, 11.5, 0.8, _fit(slide_title, 70), 30, theme["fg"], bold=True)
            if bullets:
                _textbox(slide, 0.9, HEIGHT - 1.0, 11.5, 0.6, _fit(bullets[0], 130), 16, theme["muted"])
        else:
            slide = _blank(prs, theme)
            _rect(slide, 0, 0, WIDTH, 0.12, theme["accent"])
            _textbox(slide, 0.9, 0.62, 11.6, 0.95, _fit(slide_title, 75), 32, theme["fg"], bold=True)
            _accent_rule(slide, theme)
            if bullets:
                size = 20 if len(bullets) <= 4 else (18 if len(bullets) <= 6 else 16)
                _bullets(slide, bullets, theme, size=size)
            if kind == "closing":
                _rect(slide, 0, HEIGHT - 0.9, WIDTH, 0.9, theme["panel"])
                _textbox(
                    slide, 0.9, HEIGHT - 0.68, 8.0, 0.45, "Thank you", 16, theme["accent"], bold=True
                )

        if kind not in {"title"}:
            _textbox(
                slide, WIDTH - 1.35, HEIGHT - 0.55, 0.7, 0.35,
                str(index + 1), 11, theme["muted"], align="right",
            )
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.core_properties.title = title
    prs.core_properties.author = author
    prs.core_properties.subject = subtitle or title
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))

    return {
        "path": str(path),
        "slides": len(prs.slides),
        "theme": theme_name,
        "bytes": path.stat().st_size,
        "notes": sum(1 for s in slides if s.get("notes")),
    }



def register(registry) -> None:
    @registry.tool(
        "slides.create",
        "Create a themed 16:9 PowerPoint (.pptx) with title, bullets and speaker notes.",
        schema={
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "title": {"type": "string"},
                "subtitle": {"type": "string"},
                "theme": {"type": "string", "enum": list(THEMES)},
                "author": {"type": "string"},
                "slides": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "bullets": {"type": "array", "items": {"type": "string"}},
                            "notes": {"type": "string"},
                            "kind": {
                                "type": "string",
                                "enum": ["title", "content", "image", "closing"],
                            },
                            "image": {"type": "string"},
                        },
                    },
                },
            },
            "required": ["path", "title", "slides"],
        },
        category="slides",
        requires=["python-pptx"],
    )
    def slides_create(
        ctx: ToolContext,
        path: str,
        title: str,
        slides: list,
        subtitle: str = "",
        theme: str = DEFAULT_THEME,
        author: str = "QuanAI",
    ) -> ToolResult:
        target = ctx.path(path)
        prepared = []
        for raw in slides:
            entry = dict(raw or {})
            if entry.get("image"):
                image = Path(str(entry["image"]))
                entry["image"] = str(image if image.is_absolute() else ctx.workspace / image)
            prepared.append(entry)
        info = build_deck(
            target, title, prepared, subtitle=subtitle, theme_name=theme, author=author
        )
        return ToolResult.success(
            {
                "path": ctx.rel(target),
                "slides": info["slides"],
                "theme": info["theme"],
                "notes": info["notes"],
            },
            [artifact(ctx, target, "deck", MIME, slides=info["slides"], theme=info["theme"])],
        )

    @registry.tool(
        "slides.inspect",
        "Read back a .pptx: slide count, titles and character volume.",
        schema={
            "type": "object",
            "properties": {"path": {"type": "string"}},
            "required": ["path"],
        },
        category="slides",
    )
    def slides_inspect(ctx: ToolContext, path: str) -> ToolResult:
        try:
            from pptx import Presentation  # type: ignore
        except ImportError as exc:
            raise MissingDependency(
                "slides.inspect", "python-pptx", "pip install python-pptx"
            ) from exc
        target = ctx.path(path, must_exist=True)
        deck = Presentation(str(target))
        rows = []
        for index, slide in enumerate(deck.slides, start=1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame
            ]
            body = "\n".join(texts)
            rows.append(
                {
                    "index": index,
                    "first_line": body.splitlines()[0][:100] if body else "",
                    "chars": len(body),
                    "has_notes": bool(
                        slide.has_notes_slide and slide.notes_slide.notes_text_frame.text
                    ),
                }
            )
        return ToolResult.success(
            {"path": ctx.rel(target), "slides": rows, "count": len(rows), "themes": list(THEMES)}
        )
