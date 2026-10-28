"""The CLI is the primary interface; it must work without any configuration."""

import json
from pathlib import Path

from quanai.cli import main


def _args(tmp_path: Path, *extra: str) -> list[str]:
    return ["--workspace", str(tmp_path / "ws"), "--home", str(tmp_path / "state"), *extra]


def test_tools_lists_the_surface(capsys):
    assert main(["tools"]) == 0
    out = capsys.readouterr().out
    assert "slides.create" in out and "image.generate" in out and "app.scaffold" in out


def test_doctor_works_offline(capsys):
    assert main(["doctor"]) == 0
    out = capsys.readouterr().out
    assert "Capabilities" in out
    assert "offline deterministic engine" in out


def test_plan_prints_requirements_and_tasks(tmp_path: Path, capsys):
    assert main(_args(tmp_path, "plan", "a 5 slide deck about solar energy")) == 0
    out = capsys.readouterr().out
    assert "R.deck" in out
    assert "slides.create" in out


def test_run_then_status_then_show(tmp_path: Path, capsys):
    code = main(_args(tmp_path, "run", "generate a poster image about the ocean"))
    assert code == 0
    output = capsys.readouterr().out
    assert "100.0% verified" in output

    assert main(_args(tmp_path, "status")) == 0
    listing = capsys.readouterr().out
    assert "done" in listing

    rows = [line.split() for line in listing.splitlines() if line.startswith("obj-")]
    assert rows, listing
    objective_id = rows[0][0]
    assert rows[0][1] == "done"

    assert main(_args(tmp_path, "show", objective_id)) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["id"] == objective_id


def test_run_json_output(tmp_path: Path, capsys):
    code = main(_args(tmp_path, "run", "generate an image about the ocean", "--json"))
    assert code == 0
    output = capsys.readouterr().out
    payload = json.loads(output[output.index("{") :])
    assert payload["verified"] is True
    assert payload["artifacts"]


def test_resume_reports_missing_objective(tmp_path: Path):
    try:
        main(_args(tmp_path, "resume", "obj-nope"))
    except FileNotFoundError:
        return
    raise AssertionError("expected FileNotFoundError for an unknown objective")


def test_resume_from_the_cli(tmp_path: Path, capsys):
    from quanai.config import Config
    from quanai.core.agent import Agent
    from quanai.core.task import Check, Objective, Requirement, Status, Task

    config = Config(workspace=tmp_path / "ws", home=tmp_path / "state", provider="offline")
    config.ensure_dirs()
    agent = Agent(config, provider=None)
    objective = Objective.new("resume demo")
    objective.id = "obj-cliresume"
    objective.workspace = str(config.objective_dir(objective.id))
    objective.status = Status.RUNNING
    objective.add_requirement(
        Requirement(id="R1", text="file exists", checks=[Check("file_exists", {"path": "x.txt"})])
    )
    objective.add_task(
        Task(
            id="w",
            title="write",
            tool="fs.write",
            args={"path": "x.txt", "content": "hello from resume"},
            satisfies=["R1"],
        )
    )
    agent.ledger.save(objective)

    assert main(_args(tmp_path, "resume", "obj-cliresume")) == 0
    assert (tmp_path / "ws/obj-cliresume/x.txt").exists()
    assert "100.0% verified" in capsys.readouterr().out


def test_importing_the_package_does_not_run_the_cli():
    import subprocess
    import sys

    result = subprocess.run(
        [sys.executable, "-c", "import quanai.__main__; print('imported without running')"],
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert result.returncode == 0, result.stderr
    assert "imported without running" in result.stdout
    assert "usage: quanai" not in result.stderr


def test_theme_flag_works_before_and_after_the_subcommand(tmp_path: Path, capsys):
    from pptx import Presentation

    assert main(["--workspace", str(tmp_path / "a"), "--home", str(tmp_path / "ha"),
                 "run", "a 4 slide deck about solar energy", "--theme", "sky"]) == 0
    deck = next((tmp_path / "a").glob("obj-*/deck.pptx"))
    assert str(Presentation(str(deck)).slides[0].background.fill.fore_color.rgb) == "F5F9FF"

    assert main(["--theme", "mint", "--workspace", str(tmp_path / "b"),
                 "--home", str(tmp_path / "hb"),
                 "run", "a 4 slide deck about solar energy"]) == 0
    deck = next((tmp_path / "b").glob("obj-*/deck.pptx"))
    assert str(Presentation(str(deck)).slides[0].background.fill.fore_color.rgb) == "F4FAF7"
    assert "sky" not in capsys.readouterr().out
