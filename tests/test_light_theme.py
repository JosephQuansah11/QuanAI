"""Everything QuanAI produces is light by default: decks, apps, cards, analysis art."""

import json
from pathlib import Path

import pytest

from quanai.tools import images
from quanai.tools.slides import DEFAULT_THEME, LIGHT_THEMES, THEMES
from quanai.tools.video import compose_card

DARK_CONTSTANTS = ("#0f172a", "#111c30", "#1e293b", "#38bdf8", "Brightness.dark", "color-scheme: dark")


def luminance(hex_color: str) -> float:
    value = hex_color.lstrip("#")
    red, green, blue = (int(value[i : i + 2], 16) for i in (0, 2, 4))
    return round((red * 0.299 + green * 0.587 + blue * 0.114) / 255, 3)


def png_luminance(path: Path) -> float:
    import numpy as np
    from PIL import Image

    with Image.open(path) as img:
        return float(np.asarray(img.convert("L"), dtype="float32").mean())


def test_default_theme_is_light_and_light_palettes_are_actually_light():
    assert DEFAULT_THEME in LIGHT_THEMES
    for name in LIGHT_THEMES:
        theme = THEMES[name]
        assert luminance(theme["bg"]) > 0.9, name
        assert luminance(theme["fg"]) < 0.35, name  # dark ink on light paper
        assert luminance(theme["panel"]) > 0.85, name


def contrast(a: str, b: str) -> float:
    def channel(value: int) -> float:
        ratio = value / 255
        return ratio / 12.92 if ratio <= 0.03928 else ((ratio + 0.055) / 1.055) ** 2.4

    def relative(hex_color: str) -> float:
        value = hex_color.lstrip("#")
        red, green, blue = (int(value[i : i + 2], 16) for i in (0, 2, 4))
        return 0.2126 * channel(red) + 0.7152 * channel(green) + 0.0722 * channel(blue)

    high, low = sorted((relative(a), relative(b)), reverse=True)
    return round((high + 0.05) / (low + 0.05), 2)


def test_light_themes_stay_readable():
    for name in LIGHT_THEMES:
        theme = THEMES[name]
        assert contrast(theme["fg"], theme["bg"]) >= 7.0, name
        assert contrast(theme["muted"], theme["bg"]) >= 4.5, name
        assert contrast(theme["accent"], theme["bg"]) >= 4.5, name


def test_dark_palettes_are_opt_in_only():
    assert DEFAULT_THEME not in {"midnight", "forest", "plum"}


def test_created_deck_is_light_on_every_slide(ctx, registry):
    from pptx import Presentation

    slides = [
        {"title": "Cover", "kind": "title", "subtitle": "light by default"},
        {"title": "One", "bullets": ["first point", "second point"], "notes": "say it"},
        {"title": "Two", "bullets": ["third point"], "kind": "closing"},
    ]
    result = registry.execute(
        "slides.create", {"path": "light.pptx", "title": "Light", "slides": slides}, ctx
    )
    assert result.ok, result.error
    deck = Presentation(str(Path(ctx.workspace) / "light.pptx"))
    for index, slide in enumerate(deck.slides, start=1):
        background = str(slide.background.fill.fore_color.rgb)
        assert luminance(background) > 0.9, f"slide {index} background #{background}"
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color and run.font.color.type is not None and run.text.strip():
                            assert luminance(str(run.font.color.rgb)) < 0.55, run.text


@pytest.mark.parametrize(
    "framework,platform,marker,path",
    [
        ("tkinter", "desktop", '#f7f8fa', "main.py"),
        ("pyqt6", "desktop", '#f7f8fa', "main.py"),
        ("electron", "desktop", 'background: #f7f8fa', "styles.css"),
        ("pwa", "mobile", 'color-scheme: light', "styles.css"),
        ("flutter", "mobile", 'Brightness.light', "lib/main.dart"),
        ("react-native", "mobile", "backgroundColor: '#f7f8fa'", "App.js"),
    ],
)
def test_generated_apps_are_light(ctx, registry, framework, platform, marker, path):
    result = registry.execute(
        "app.scaffold",
        {"path": f"apps/{framework}", "description": "a habit tracker", "platform": platform, "framework": framework},
        ctx,
    )
    assert result.ok, result.error
    root = Path(ctx.workspace) / f"apps/{framework}"
    source = (root / path).read_text()
    assert marker in source
    for dark in DARK_CONTSTANTS:
        assert dark not in source, f"{framework}/{path} still contains {dark}"


def test_pwa_manifest_and_icons_are_light(ctx, registry):
    registry.execute(
        "app.scaffold",
        {"path": "pwa", "description": "a notes app", "platform": "mobile", "framework": "pwa"},
        ctx,
    )
    root = Path(ctx.workspace) / "pwa"
    manifest = json.loads((root / "manifest.webmanifest").read_text())
    assert luminance(manifest["background_color"]) > 0.9
    assert manifest["theme_color"].lower() == "#2563eb"
    for size in (192, 512):
        assert png_luminance(root / f"icons/icon-{size}.png") > 200


def test_storyboard_cards_and_video_frames_are_light(tmp_path):
    background = tmp_path / "bg.png"
    images.render_procedural("abstract night city", 640, 360, seed=3).save(background)

    flat = compose_card(None, "A caption on a light card", "with a subtitle", 640, 360)
    flat_path = tmp_path / "flat.png"
    flat.save(flat_path)
    assert png_luminance(flat_path) > 200

    over_art = compose_card(background, "A caption over artwork", "with a subtitle", 640, 360)
    over_path = tmp_path / "over.png"
    over_art.save(over_path)
    assert png_luminance(over_path) > 160


def test_analysis_visuals_are_light(ctx, registry):
    from pathlib import Path as _Path

    from quanai.util import ffmpeg_bin, run_cmd

    if not ffmpeg_bin():
        pytest.skip("ffmpeg unavailable")
    workspace = _Path(ctx.workspace)
    (workspace / "media").mkdir(parents=True, exist_ok=True)
    audio = workspace / "media/tone.wav"
    run_cmd(
        [ffmpeg_bin(), "-y", "-loglevel", "error", "-f", "lavfi",
         "-i", "sine=frequency=440:duration=2", "-ac", "1", str(audio)],
        timeout=60,
    )
    result = registry.execute(
        "media.ingest", {"path": "media/tone.wav", "out_dir": "media/out", "transcribe": False}, ctx
    )
    assert result.ok, result.error
    out = workspace / "media/out"
    assert png_luminance(out / "audio_waveform.png") > 200
    assert png_luminance(out / "audio_spectrogram.png") > 200


def test_procedural_art_is_light_but_still_structured(tmp_path):
    for style in ("editorial", "minimal", "technical", "organic", "luxury"):
        target = tmp_path / f"art-{style}.png"
        images.render_procedural("abstract ocean poster", 480, 270, seed=2, style=style).save(target)
        mean = png_luminance(target)
        assert mean > 180, f"{style} is not light ({mean:.0f})"

        import numpy as np
        from PIL import Image

        with Image.open(target) as img:
            std = float(np.asarray(img.convert("L"), dtype="float32").std())
        assert std > 12, f"{style} lost all structure (std {std:.1f})"


def test_dark_art_remains_available_on_request(tmp_path):
    target = tmp_path / "dark.png"
    images.render_procedural("abstract ocean poster", 320, 180, seed=1, light=False).save(target)
    assert png_luminance(target) < 140


def test_image_tool_defaults_to_light_and_can_opt_out(ctx, registry):
    light = registry.execute(
        "image.generate", {"path": "a.png", "prompt": "deep ocean", "width": 320, "height": 180}, ctx
    )
    assert light.ok, light.error
    assert light.data["light"] is True
    assert png_luminance(Path(ctx.workspace) / "a.png") > 180

    dark = registry.execute(
        "image.generate",
        {"path": "b.png", "prompt": "deep ocean", "width": 320, "height": 180, "light": False},
        ctx,
    )
    assert dark.ok, dark.error
    assert dark.data["light"] is False
    assert png_luminance(Path(ctx.workspace) / "b.png") < 140
