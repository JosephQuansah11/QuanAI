"""Every tool must produce a real, inspectable artifact."""

import json
from pathlib import Path

import pytest

from quanai.util import ffmpeg_bin


HAS_FFMPEG = bool(ffmpeg_bin())


def test_fs_write_read_and_list(ctx, registry):
    written = registry.execute("fs.write", {"path": "notes/a.txt", "content": "hello\nworld"}, ctx)
    assert written.ok, written.error
    assert written.artifacts[0].bytes == 11

    read = registry.execute("fs.read", {"path": "notes/a.txt"}, ctx)
    assert read.data["content"] == "hello\nworld"

    listed = registry.execute("fs.list", {"path": "."}, ctx)
    assert any(row["path"] == "notes/a.txt" for row in listed.data["files"])


def test_path_escape_is_refused(ctx, registry):
    result = registry.execute("fs.write", {"path": "../../escape.txt", "content": "nope"}, ctx)
    assert not result.ok
    assert "escapes" in result.error.lower()


def test_slides_create_produces_openable_deck(ctx, registry):
    from pptx import Presentation

    slides = [
        {"title": "Cover", "kind": "title"},
        {"title": "Why it matters", "bullets": ["point one", "point two"], "notes": "say this"},
        {"title": "Next", "bullets": ["do the thing"], "kind": "closing"},
    ]
    result = registry.execute(
        "slides.create",
        {"path": "decks/demo.pptx", "title": "Demo", "subtitle": "sub", "slides": slides},
        ctx,
    )
    assert result.ok, result.error
    target = Path(ctx.workspace) / "decks/demo.pptx"
    deck = Presentation(str(target))
    assert len(deck.slides) == 3
    assert deck.slides[1].notes_slide.notes_text_frame.text == "say this"
    assert result.artifacts[0].kind == "deck"


def test_image_generate_is_deterministic_and_not_blank(ctx, registry):
    from PIL import Image
    import numpy as np

    first = registry.execute(
        "image.generate", {"path": "art/one.png", "prompt": "solar city", "width": 320, "height": 240}, ctx
    )
    second = registry.execute(
        "image.generate", {"path": "art/two.png", "prompt": "solar city", "width": 320, "height": 240}, ctx
    )
    assert first.ok and second.ok
    assert first.data["backend"] == "procedural"
    a = np.asarray(Image.open(Path(ctx.workspace) / "art/one.png").convert("L"))
    b = np.asarray(Image.open(Path(ctx.workspace) / "art/two.png").convert("L"))
    assert a.std() > 10  # real structure, not a flat fill
    assert np.array_equal(a, b)  # same prompt + seed -> same pixels


def test_image_variations_differ(ctx, registry):
    result = registry.execute(
        "image.variations",
        {"path_prefix": "art/var", "prompt": "deep ocean", "count": 2, "width": 256, "height": 144},
        ctx,
    )
    assert result.ok, result.error
    assert result.data["count"] == 2
    sizes = {art.bytes for art in result.artifacts}
    assert all(size > 1000 for size in sizes)


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg not available")
def test_video_from_images_records_real_duration(ctx, registry):
    for index in range(2):
        registry.execute(
            "image.generate",
            {"path": f"frames/f{index}.png", "prompt": f"scene {index}", "width": 320, "height": 180},
            ctx,
        )
    result = registry.execute(
        "video.from_images",
        {"path": "out/clip.mp4", "images": ["frames/f0.png", "frames/f1.png"], "seconds_per_image": 1.0},
        ctx,
    )
    assert result.ok, result.error
    assert result.data["duration"] >= 1.5
    assert (Path(ctx.workspace) / "out/clip.mp4").stat().st_size > 1000


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg not available")
def test_storyboard_then_video(ctx, registry):
    scenes = [
        {"caption": "First beat", "image_prompt": "abstract sunrise", "seconds": 1.0, "seed": 1},
        {"caption": "Second beat", "image_prompt": "abstract dusk", "seconds": 1.0, "seed": 2},
    ]
    board = registry.execute(
        "video.storyboard",
        {"out_dir": "story", "title": "T", "scenes": scenes, "width": 480, "height": 270},
        ctx,
    )
    assert board.ok, board.error
    assert board.data["count"] == 2

    video = registry.execute(
        "video.from_images",
        {"path": "story.mp4", "images": board.data["cards"], "seconds_per_image": 1.0, "width": 480, "height": 270},
        ctx,
    )
    assert video.ok, video.error
    assert video.data["duration"] >= 1.5


def test_app_scaffold_python_framework_compiles(ctx, registry):
    result = registry.execute(
        "app.scaffold",
        {"path": "app", "description": "a habit tracker for daily routines", "framework": "tkinter"},
        ctx,
    )
    assert result.ok, result.error
    source = (Path(ctx.workspace) / "app/main.py").read_text()
    compile(source, "main.py", "exec")
    assert "tkinter" in source and "class App" in source
    assert (Path(ctx.workspace) / "app/README.md").exists()


def test_app_scaffold_web_frameworks_produce_valid_assets(ctx, registry):
    for framework, platform in (("pwa", "mobile"), ("electron", "desktop"), ("react-native", "mobile")):
        result = registry.execute(
            "app.scaffold",
            {"path": f"apps/{framework}", "description": "a notes app", "platform": platform, "framework": framework},
            ctx,
        )
        assert result.ok, result.error
        root = Path(ctx.workspace) / f"apps/{framework}"
        assert (root / "README.md").exists()
        if framework == "pwa":
            manifest = json.loads((root / "manifest.webmanifest").read_text())
            assert manifest["display"] == "standalone"
            assert (root / "icons/icon-192.png").exists()
        if framework in {"electron", "react-native"}:
            package = json.loads((root / "package.json").read_text())
            assert package["name"]


def test_code_zip_and_report(ctx, registry):
    registry.execute("fs.write", {"path": "proj/main.py", "content": "print('hi')\n"}, ctx)
    zipped = registry.execute("code.zip", {"path": "proj", "archive": "proj.zip"}, ctx)
    assert zipped.ok and zipped.data["members"] == 1

    report = registry.execute(
        "report.write",
        {"path": "doc.md", "title": "Solar briefing", "request": "explain solar energy in cities"},
        ctx,
    )
    assert report.ok
    text = (Path(ctx.workspace) / "doc.md").read_text()
    assert text.startswith("# Solar briefing")
    assert text.count("## ") >= 3


def test_plan_tools_write_json(ctx, registry):
    deck = registry.execute("plan.deck", {"request": "a deck about urban farming", "slides": 5}, ctx)
    assert deck.ok, deck.error
    assert deck.data["count"] == 5
    payload = json.loads((Path(ctx.workspace) / "plans/deck.json").read_text())
    assert payload["title"] and len(payload["slides"]) == 5

    video = registry.execute("plan.video", {"request": "a promo video about farming", "scenes": 3}, ctx)
    assert video.ok and len(video.data["scenes"]) == 3

    app = registry.execute("plan.app", {"request": "a farm log app", "platform": "desktop"}, ctx)
    assert app.ok and app.data["name"]


def test_shell_is_disabled_by_default(ctx, registry):
    result = registry.execute("code.shell", {"command": "echo hi"}, ctx)
    assert not result.ok
    assert "disabled" in result.error


def test_python_run_executes_inside_workspace(ctx, registry):
    registry.execute("fs.write", {"path": "scripts/hi.py", "content": "print('from script')\n"}, ctx)
    result = registry.execute("code.python_run", {"path": "scripts/hi.py"}, ctx)
    assert result.ok, result.error
    assert "from script" in result.data["stdout"]


def test_missing_required_argument_is_reported(ctx, registry):
    result = registry.execute("image.generate", {"path": "x.png"}, ctx)
    assert not result.ok
    assert "prompt" in result.error
