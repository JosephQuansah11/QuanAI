"""Verification checks: they must measure, not trust."""

from pathlib import Path

import pytest

from quanai.core.task import Check, Evidence, Objective, Requirement, Status
from quanai.core.verify import Verifier, VerifyContext, known_checks
from quanai.util import ffmpeg_bin

HAS_FFMPEG = bool(ffmpeg_bin())

MAIN_PY = "import tkinter\n\n\ndef main():\n    print('hello from the app')\n"


@pytest.fixture()
def verifier(tmp_path: Path) -> Verifier:
    workspace = tmp_path / "ws"
    workspace.mkdir()
    return Verifier(workspace)


@pytest.fixture()
def vctx(tmp_path: Path) -> VerifyContext:
    workspace = tmp_path / "ws"
    workspace.mkdir(exist_ok=True)
    return VerifyContext(workspace=workspace)


def test_known_checks_cover_the_core_artifact_types():
    kinds = set(known_checks())
    assert {
        "file_exists", "file_min_bytes", "text_contains", "json_valid", "zip_valid",
        "pptx_valid", "pptx_notes", "image_valid", "video_valid", "audio_valid",
        "python_compiles", "js_syntax", "scaffold_files", "markdown_report", "media_report",
    } <= kinds


def test_missing_file_fails_with_a_reason(verifier, vctx):
    evidence = verifier.run_check(Check("file_exists", {"path": "nope.txt"}), vctx)
    assert evidence.ok is False
    assert "missing file" in evidence.detail


def test_text_contains_reports_the_missing_phrase(verifier, vctx):
    (Path(vctx.workspace) / "a.md").write_text("hello world")
    assert verifier.run_check(Check("text_contains", {"path": "a.md", "needles": ["hello"]}), vctx).ok
    bad = verifier.run_check(
        Check("text_contains", {"path": "a.md", "needles": ["goodbye"], "min_chars": 999}), vctx
    )
    assert not bad.ok and "goodbye" in bad.detail


def test_json_valid_checks_keys(verifier, vctx):
    (Path(vctx.workspace) / "a.json").write_text('{"x": 1}')
    assert verifier.run_check(Check("json_valid", {"path": "a.json", "keys": ["x"]}), vctx).ok
    assert not verifier.run_check(Check("json_valid", {"path": "a.json", "keys": ["y"]}), vctx).ok


def test_pptx_valid_measures_slides_and_terms(verifier, vctx):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for title in ("Solar cover", "Costs fall fast", "Grid capacity"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"{title} - solar energy adoption in cities"
    deck.save(str(Path(vctx.workspace) / "d.pptx"))

    ok = verifier.run_check(
        Check("pptx_valid", {"path": "d.pptx", "min_slides": 3, "required_terms": ["solar"]}), vctx
    )
    assert ok.ok, ok.detail
    assert ok.observed["slides"] == 3

    too_many = verifier.run_check(Check("pptx_valid", {"path": "d.pptx", "min_slides": 10}), vctx)
    assert not too_many.ok and "3 slide" in too_many.detail


def test_image_valid_rejects_a_blank_image(verifier, vctx):
    import numpy as np
    from PIL import Image

    Image.new("RGB", (200, 120), (12, 12, 12)).save(Path(vctx.workspace) / "flat.png")
    evidence = verifier.run_check(Check("image_valid", {"path": "flat.png"}), vctx)
    assert not evidence.ok
    assert "blank" in evidence.detail

    noise = (np.random.default_rng(0).random((120, 200, 3)) * 255).astype("uint8")
    Image.fromarray(noise).save(Path(vctx.workspace) / "noise.png")
    assert verifier.run_check(Check("image_valid", {"path": "noise.png"}), vctx).ok


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg not available")
def test_video_and_audio_checks_measure_duration(verifier, vctx):
    from quanai.util import run_cmd

    target = Path(vctx.workspace)
    (target / "media").mkdir(exist_ok=True)
    run_cmd(
        [ffmpeg_bin(), "-y", "-loglevel", "error", "-f", "lavfi",
         "-i", "testsrc=size=160x120:rate=10:duration=2", "-pix_fmt", "yuv420p",
         str(target / "media/v.mp4")],
        timeout=90,
    )
    run_cmd(
        [ffmpeg_bin(), "-y", "-loglevel", "error", "-f", "lavfi",
         "-i", "sine=frequency=440:duration=2", "-ac", "1", str(target / "media/a.wav")],
        timeout=90,
    )
    video = verifier.run_check(Check("video_valid", {"path": "media/v.mp4", "min_seconds": 1.5}), vctx)
    assert video.ok, video.detail
    audio = verifier.run_check(Check("audio_valid", {"path": "media/a.wav", "min_seconds": 1.5}), vctx)
    assert audio.ok, audio.detail


def test_scaffold_and_compile_checks(verifier, vctx):
    root = Path(vctx.workspace) / "app"
    root.mkdir()
    (root / "main.py").write_text(MAIN_PY)
    (root / "README.md").write_text("# app\n\nA small application.\n")
    assert verifier.run_check(
        Check("scaffold_files", {"path": "app", "required": ["main.py", "README.md"]}), vctx
    ).ok
    assert verifier.run_check(Check("python_compiles", {"path": "app"}), vctx).ok
    assert not verifier.run_check(
        Check("scaffold_files", {"path": "app", "required": ["missing.py"]}), vctx
    ).ok

    (root / "broken.py").write_text("def oops(:\n  pass\n")
    broken = verifier.run_check(Check("python_compiles", {"path": "app"}), vctx)
    assert not broken.ok and "syntax" in broken.detail.lower()


def test_unknown_check_raises(verifier, vctx):
    from quanai.errors import CheckError

    with pytest.raises(CheckError):
        verifier.run_check(Check("does_not_exist", {}), vctx)


def test_verify_objective_records_evidence(tmp_path: Path):
    workspace = tmp_path / "ws2"
    workspace.mkdir()
    (workspace / "present.txt").write_text("here")
    objective = Objective.new("verify me", workspace=str(workspace))
    objective.add_requirement(
        Requirement(id="R1", text="present", checks=[Check("file_exists", {"path": "present.txt"})])
    )
    objective.add_requirement(
        Requirement(id="R2", text="absent", checks=[Check("file_exists", {"path": "gone.txt"})])
    )
    verification = Verifier(workspace).verify_objective(objective)
    assert verification["passed"] == 1
    assert verification["failed"] == 1
    assert verification["results"] == {"R1": True, "R2": False}
    assert objective.requirement("R1").status is Status.DONE
    assert objective.requirement("R2").status is Status.FAILED
    assert all(isinstance(ev, Evidence) for ev in objective.requirement("R1").evidence)


def test_broken_check_does_not_crash_the_run(verifier, vctx, monkeypatch):
    import quanai.core.verify as verify_module

    def explode(ctx, **kwargs):
        raise RuntimeError("boom")

    monkeypatch.setitem(verify_module._REGISTRY, "explode_check", explode)
    evidence = verifier.run_check(Check("explode_check", {}), vctx)
    assert evidence.ok is False
    assert "boom" in evidence.detail
